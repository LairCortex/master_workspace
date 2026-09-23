"""XLSX import against the active game calendar (piece C3a, tasks 4.1–4.2).

Spec «Колонки листа» (xlsx-import): parsed numbers become coordinates of the
active calendar — an absent day (31 in a 30-day game month) imports as the
last valid day of the same month, a month beyond the month count as the last
month's last day, and every such transfer lands in the final report ("перенесённые
даты", its own section per spec «Итоговый отчёт импорта») with sheet, row number,
field and old → new dates; valid dates never get a transfer row, so under the
standard preset the section is empty.  Auto-created ghost link targets receive
the already-shifted dates (design D7: shared route before recording), and the
write itself goes through the routed resolver: coordinate columns under a
custom calendar, date columns under the preset (design D3).
"""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import Workbook

from app.application.services.xlsx_import_service import _coord_text, XlsxImportService
from app.domain.game_calendar import (
    CalendarSpec,
    CustomCalendar,
    IntercalaryDay,
    IntercalarySpec,
    MonthDay,
    MonthSpec,
    reset_current_calendar,
    set_current_calendar,
)
from app.infrastructure.calendar_storage import (
    encode_coord,
)
from app.infrastructure.db.uow import GameSessionUoW
from app.infrastructure.db.models import CharacterModel, LocationModel, resolve_coord


# Ten months of thirty days: ISO 31-е overflows any month, months 11–12 are
# out of count — the two shift fixtures of design D4 in the import route.
_SPEC = CalendarSpec(
    months=tuple(
        MonthSpec(name, 30)
        for name in (
            "Медвежарь", "Ледокол", "Травень", "Цветень", "Жневень",
            "Сенокос", "Гридень", "Листопад", "Хмурень", "Студень",
        )
    ),
    week_names=("пн", "вт", "ср", "чт", "пт", "сб", "вс"),
)
_CUSTOM = CustomCalendar(_SPEC)

# Ten months of forty-five days: a day the clamp can land on that no Gregorian
# month has — the transfer row then signs its new date with the codec string
# instead of a fabricated ISO (the Iso rule of design D5, same as iso_or_coord).
_WIDE_SPEC = CalendarSpec(
    months=tuple(
        MonthSpec(name, 45)
        for name in (
            "Медвежарь", "Ледокол", "Травень", "Цветень", "Жневень",
            "Сенокос", "Гридень", "Листопад", "Хмурень", "Студень",
        )
    ),
    week_names=("пн", "вт", "ср", "чт", "пт", "сб", "вс"),
)
_WIDE = CustomCalendar(_WIDE_SPEC)

# C5 task 2.2/2.3 — the same ten-month layout plus one declared intercalary
# day: the game wording then routes an IntercalaryDay coordinate through the
# exact shift/report pipeline the ISO forms already travel.
_INTER_SPEC = CalendarSpec(
    months=tuple(
        MonthSpec(name, 30)
        for name in (
            "Медвежарь", "Ледокол", "Травень", "Цветень", "Жневень",
            "Сенокос", "Гридень", "Листопад", "Хмурень", "Студень",
        )
    ),
    week_names=("пн", "вт", "ср", "чт", "пт", "сб", "вс"),
    intercalary=(IntercalarySpec("Медожор", after_month=1),),
)
_INTER = CustomCalendar(_INTER_SPEC)


@pytest.fixture(autouse=True)
def _standard_active_calendar():
    """An imported global must not leak in or out of any test here."""
    reset_current_calendar()
    yield
    reset_current_calendar()


def _svc() -> XlsxImportService:
    return XlsxImportService()


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _sheet(wb: Workbook, title: str, headers, rows) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)


def _save(tmp_path, wb: Workbook, name: str = "import.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return path


FULL_HEADERS = ["Имя", "Дата начала", "Дата конца"]
EVENT_HEADERS = ["Имя", "Дата начала", "Связь локациями"]


async def _one(session, model, **filters):
    from sqlalchemy import select

    stmt = select(model)
    for attr, value in filters.items():
        stmt = stmt.where(getattr(model, attr) == value)
    return (await session.execute(stmt)).scalars().one()


# ── 4.1 — coordinates, shift policy, transfer rows ────────────────────────

class TestParsedDatesBecomeCoordinates:
    async def test_iso_day_31_imports_as_30_with_report_row(self, tmp_path, async_session):
        # Spec scenario «Дня нет в игровом календаре — импорт переносит».
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Иван", "2026-08-31", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Иван")
        # Custom route: the coordinate columns carry the clamped day; the
        # NOT NULL legacy date column only holds the schema placeholder the
        # routed write leaves for a fresh coordinate row (same rule as the
        # repositories' CoordMappingMixin — nothing reads it while coord wins).
        assert char.start_coord == encode_coord(MonthDay(2026, 8, 30))
        assert char.start_date_raw == date(1, 1, 1)  # слот-заглушка NOT NULL
        assert resolve_coord(char, "start") == MonthDay(2026, 8, 30)

        shift, = report.date_shifts
        assert (shift.sheet, shift.row_number, shift.field) == (
            "Персонажи", 2, "Дата начала",
        )
        assert (shift.old, shift.new) == ("2026-08-31", "2026-08-30")

    async def test_month_beyond_count_imports_last_month_last_day(
        self, tmp_path, async_session
    ):
        # Design D4/month-out-of-count clamp: last existing month, its last day.
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Марья", "2026-11-05", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Марья")
        assert char.start_coord == encode_coord(MonthDay(2026, 10, 30))

        shift, = report.date_shifts
        assert shift.field == "Дата начала"
        assert (shift.old, shift.new) == ("2026-11-05", "2026-10-30")

    async def test_valid_dates_have_no_transfer_rows(self, tmp_path, async_session):
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Фёдор", "2026-08-15", "2026-09-02"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Фёдор")
        assert char.start_coord == encode_coord(MonthDay(2026, 8, 15))
        assert char.end_coord == encode_coord(MonthDay(2026, 9, 2))
        assert report.date_shifts == []

    async def test_bc_date_shift_keeps_the_era(self, tmp_path, async_session):
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        # Signed-ISO BC form parses to the same coordinate the text form would.
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Ксеркс", "-0044-08-31", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Ксеркс")
        assert char.start_coord == encode_coord(MonthDay(44, 8, 30))
        assert bool(char.start_bc) is True  # era flag is an INTEGER column

        shift, = report.date_shifts
        assert (shift.old, shift.new) == ("0044-08-31 до н.э.", "0044-08-30 до н.э.")

    async def test_standard_calendar_imports_iso_into_date_columns(
        self, tmp_path, async_session
    ):
        # Bit-identical standard behavior: preset active ⇒ date columns, no
        # coordinate columns, no transfer rows, no shift.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Обычный", "2026-08-31", "2026-09-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Обычный")
        assert (char.start_date, char.end_date) == (date(2026, 8, 31), date(2026, 9, 1))
        assert char.start_coord is None and char.end_coord is None
        assert report.date_shifts == []

    async def test_ghost_targets_receive_the_shifted_dates(self, tmp_path, async_session):
        # Design D7: the shared route shifts before recording, so an
        # auto-created link target is born with the clamped dates.
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "2026-08-31", "Парк"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        # the report is not asserted here — the ghost target dates are checked below
        await _svc().apply_plan(plan, GameSessionUoW(async_session))

        ghost = plan.ghosts[("location", "парк")]
        assert ghost.min_start == MonthDay(2026, 8, 30)

        park = await _one(async_session, LocationModel, name="Парк")
        assert park.start_coord == encode_coord(MonthDay(2026, 8, 30))
        assert park.end_coord == encode_coord(MonthDay(2026, 8, 30))
        assert park.start_date_raw == date(1, 1, 1)  # fresh coord row → placeholder

    async def test_merge_shift_row_follows_the_winning_contribution(
        self, tmp_path, async_session
    ):
        # The later row replaces the field wholesale (merge rule), so its
        # transfer state replaces the earlier one's too — a superseded
        # transfer never reaches the report ("only фактические переносы").
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS,
               [["Иван", "2026-08-31", None], ["Иван", "2026-07-05", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Иван")
        assert char.start_coord == encode_coord(MonthDay(2026, 7, 5))
        assert report.date_shifts == []

    async def test_merge_shift_follows_a_later_row_that_shifts(
        self, tmp_path, async_session
    ):
        # The mirror case: the earlier date needed no transfer, the winning
        # later one does — the report lists the winner's transfer (the row
        # number is the winning contribution's line).
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS,
               [["Иван", "2026-07-05", None], ["Иван", "2026-08-31", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Иван")
        assert char.start_coord == encode_coord(MonthDay(2026, 8, 30))
        shift, = report.date_shifts
        assert (shift.row_number, shift.old, shift.new) == (3, "2026-08-31", "2026-08-30")

    async def test_skipped_row_transfer_never_reaches_the_report(
        self, tmp_path, async_session
    ):
        # An ambiguous DB reference skips the referring row — its shifted date
        # is never imported, so the section must not list it.
        set_current_calendar(_CUSTOM)
        from app.infrastructure.db.models import LocationModel as Loc

        async_session.add(Loc(name="Парк", start_date=date(2000, 1, 1)))
        async_session.add(Loc(name="парк", start_date=date(2000, 1, 2)))
        await async_session.flush()

        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "2026-08-31", "Парк"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        assert report.skipped and report.created == 0
        assert report.date_shifts == []

    async def test_game_form_day_beyond_month_shifts_like_iso(
        self, tmp_path, async_session
    ):
        # C5 task 2.1–2.3, spec «Игровая дата с числом длиннее месяца
        # переносится»: the game-form coordinate 31 (Листопад is month 8)
        # clamps exactly like the ISO 31-е did, with the usual transfer row.
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Листопалый", "31 Листопад 2026", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Листопалый")
        assert resolve_coord(char, "start") == MonthDay(2026, 8, 30)

        shift, = report.date_shifts
        assert (shift.sheet, shift.row_number, shift.field) == (
            "Персонажи", 2, "Дата начала",
        )
        assert (shift.old, shift.new) == ("2026-08-31", "2026-08-30")

    async def test_named_intercalary_day_imports_as_intercalary_coordinate(
        self, tmp_path, async_session
    ):
        # C5 task 2.2/2.3: a name declared in spec.intercalary lands the
        # same-calendar IntercalaryDay coordinate (rule index 0) over the
        # routed write — codec coordinate, era beside it, and no transfer.
        set_current_calendar(_INTER)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS,
               [["Медожин", "Медожор 44 г. до н.э.", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Медожин")
        assert char.start_coord == encode_coord(IntercalaryDay(44, 0))
        assert bool(char.start_bc) is True
        assert resolve_coord(char, "start") == IntercalaryDay(44, 0)
        assert report.date_shifts == []


# ── 4.2 — the report's own "перенесённые даты" section ────────────────────

class TestReportDateShiftSection:
    async def test_report_lists_only_actual_shifts(self, tmp_path, async_session):
        # Spec scenario «Переносы видимы в отчёте»: shifted + untouched rows,
        # the section names exactly the фактически сдвинутые даты.
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [
            ["Валидный", "2026-05-10", "2026-06-10"],   # no rows
            ["Перенесён", "2026-08-31", None],          # start row
            ["Оба", "-0044-07-31", "-0044-08-31"],      # start + end rows, BC
        ])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        assert report.created == 3  # all rows imported (the shifts are silent-free)
        assert [(s.sheet, s.row_number, s.field, s.old, s.new) for s in report.date_shifts] == [
            ("Персонажи", 3, "Дата начала", "2026-08-31", "2026-08-30"),
            ("Персонажи", 4, "Дата начала", "0044-07-31 до н.э.", "0044-07-30 до н.э."),
            ("Персонажи", 4, "Дата конца", "0044-08-31 до н.э.", "0044-08-30 до н.э."),
        ]

    async def test_standard_calendar_section_stays_empty(self, tmp_path, async_session):
        # Spec task 4.2: under the standard preset there is no section at all.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["А", "2026-08-31", "2026-09-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))
        assert report.date_shifts == []


# ── the transfer row's date signature (design D5's Iso rule) ───────────────

class TestTransferDateSignature:
    async def test_transfer_to_a_non_gregorian_day_is_signed_with_the_codec(
        self, tmp_path, async_session
    ):
        # The 45th day of the last game month is a real game coordinate and no
        # Gregorian date — the report prints the codec text, never a fabricated
        # ISO (the same Iso rule iso_or_coord follows for display carriers).
        set_current_calendar(_WIDE)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Широкий", "2026-11-05", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        report = await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Широкий")
        assert char.start_coord == encode_coord(MonthDay(2026, 10, 45))
        shift, = report.date_shifts
        assert (shift.old, shift.new) == ("2026-11-05", "M:2026:10:45")

    def test_intercalary_coordinate_is_signed_with_the_codec(self):
        # The signer's generic coordinate contract: a day outside the months
        # has no ISO form at all, its codec text stands in, era marker kept.
        assert _coord_text(IntercalaryDay(44, 0), False) == "I:44:0"
        assert _coord_text(IntercalaryDay(44, 0), True) == "I:44:0 до н.э."


# ── A3: one shared route also for the importer — value collisions ───────────

class TestPlaceholderValueCollision:
    """The insert placeholder (date(1, 1, 1)) is schema noise for the NOT NULL
    legacy slot; the user may import that very date as real data. The shared
    ``route_insert_dates`` must keep the two roles apart on both storages.
    """

    async def test_custom_route_placeholder_never_outranks_the_real_coord(
        self, tmp_path, async_session
    ):
        # The imported coordinate's date equals the placeholder: the slot gets
        # the same literal as schema noise, yet the read route answers the
        # coordinate — the placeholder is never mistaken for the value.
        set_current_calendar(_CUSTOM)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Нулевой", "0001-01-01", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Нулевой")
        assert char.start_coord == encode_coord(MonthDay(1, 1, 1))
        assert char.start_date_raw == date(1, 1, 1)  # placeholder == value, by design
        assert resolve_coord(char, "start") == MonthDay(1, 1, 1)

    async def test_standard_route_stores_the_colliding_date_as_real_data(
        self, tmp_path, async_session
    ):
        # Under the preset the same literal is an ordinary value written to
        # the date columns: no coordinate slot, no placeholder distinction —
        # the shared route only placeholders rows whose value really went to
        # the coordinate columns.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", FULL_HEADERS, [["Нулевой", "0001-01-01", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        await _svc().apply_plan(plan, GameSessionUoW(async_session))

        char = await _one(async_session, CharacterModel, name="Нулевой")
        assert char.start_date_raw == date(1, 1, 1)  # the real imported value
        assert char.start_coord is None
        assert resolve_coord(char, "start") == MonthDay(1, 1, 1)
