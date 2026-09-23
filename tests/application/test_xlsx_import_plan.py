"""Tests for the import-plan pre-analysis (task group 2, rework-xlsx-import).

Covers ``XlsxImportService.analyze_file``:
* 2.1 — workbook-level branch matrix: only-.xlsx reading, case-insensitive
  known-sheet mapping, unknown-sheet warnings, every fatal branch
  (missing/corrupted/empty file, no known sheets, broken known sheet);
* 2.2 — in-file merge by (type, lower(name)) with later-non-empty override
  and link accumulation, planned skips (empty name / broken start date /
  out-of-range rating) with sheet + row number;
* 2.3 — name index against an in-memory DB with duplicate names: unique
  match → update plan (non-empty fields only), ambiguous row name → create
  plan, ambiguous link reference → planned skip; ghosts for unknown targets
  with min-start/max-end date buffers.
"""
from datetime import date

from openpyxl import Workbook
from sqlalchemy import func, select

from app.application.services.xlsx_import_service import (
    LINK_TO_DB,
    LINK_TO_FILE,
    LINK_TO_GHOST,
    XlsxImportService,
)
from app.domain.game_calendar import MonthDay
from app.infrastructure.db.uow import GameSessionUoW
from app.infrastructure.db.models import CharacterModel, ItemModel, OrganizationModel


def _svc() -> XlsxImportService:
    # The service has no per-type service dependency (single-sheet API removed).
    return XlsxImportService()


def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default "Sheet"
    return wb


def _sheet(wb: Workbook, title: str, headers, rows) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)


def _save(tmp_path, wb: Workbook, name: str = "plan.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return path


CHAR_HEADERS = ["Имя", "Дата начала"]
EVENT_HEADERS = ["Имя", "Дата начала", "Связь персонажами"]


# ── 2.1 — fatal file-level branches ───────────────────────────────────────

class TestFatalBranches:
    async def test_missing_file_is_fatal(self, tmp_path):
        plan = await _svc().analyze_file(tmp_path / "nope.xlsx")
        assert plan.has_fatal and plan.entities == {}
        assert plan.fatal_errors == [f"Файл не найден: {tmp_path / 'nope.xlsx'}"]

    async def test_corrupted_file_is_fatal(self, tmp_path):
        p = tmp_path / "bad.xlsx"
        p.write_bytes(b"not a zip archive")
        plan = await _svc().analyze_file(p)
        assert plan.has_fatal
        assert "повреждён" in plan.fatal_errors[0]
        assert not plan.skipped_rows

    async def test_only_unknown_sheets_is_fatal_with_warning(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Заметки", ["Что-то"], [["..."]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "знакомого" in plan.fatal_errors[0]
        # The unknown sheet is additionally listed as a warning.
        assert any("Заметки" in w for w in plan.warnings)

    async def test_known_sheets_with_headers_only_are_fatal(self, tmp_path):
        # Spec scenario «Пустой файл»: знакомые листы есть, строк данных нет.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [])
        _sheet(wb, "События", EVENT_HEADERS, [])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "пустой" in plan.fatal_errors[0]
        assert plan.entities == {} and not plan.skipped_rows

    async def test_default_empty_workbook_is_fatal(self, tmp_path):
        # Workbook() alone (single default "Sheet") — no known sheets at all.
        p = tmp_path / "empty.xlsx"
        Workbook().save(p)
        plan = await _svc().analyze_file(p)
        assert plan.has_fatal and "знакомого" in plan.fatal_errors[0]

    async def test_known_sheet_without_header_row_is_fatal(self, tmp_path):
        wb = _new_workbook()
        ws = wb.create_sheet("События")
        ws.append([None, None])  # "blank" first row — no header row
        ws.append(["Е", date(2001, 1, 1)])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "События" in plan.fatal_errors[0]
        assert "заголов" in plan.fatal_errors[0]

    async def test_known_sheet_missing_required_column_is_fatal(self, tmp_path):
        # Spec scenario: «Персонажи» without «Имя» — import must not start.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Дата начала"], [[date(2001, 1, 1)]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "Персонажи" in plan.fatal_errors[0]
        assert "Имя" in plan.fatal_errors[0]
        assert plan.entities == {}

    async def test_every_broken_known_sheet_reported(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", ["Дата начала"], [[date(2001, 1, 1)]])       # no Имя
        _sheet(wb, "Предметы", ["Имя"], [["x"]])                           # no Дата начала
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert len(plan.fatal_errors) == 2
        assert any("События" in e for e in plan.fatal_errors)
        assert any("Предметы" in e and "Дата начала" in e for e in plan.fatal_errors)

    async def test_duplicate_sheet_name_variants_are_fatal(self, tmp_path):
        # Excel/openpyxl refuse exact case-duplicates but a whitespace variant
        # ("Персонажи " vs "Персонажи") is representable and normalizes to the
        # same registry sheet — the parser cannot pick one, so it is fatal.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["А", date(2001, 1, 1)]])
        _sheet(wb, "Персонажи ", CHAR_HEADERS, [["Б", date(2001, 1, 1)]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "Дубликаты" in plan.fatal_errors[0]

    async def test_all_rows_problematic_is_not_fatal_but_plans_nothing(self, tmp_path):
        # Row problems are planned skips (spec separates them from file-level
        # fatal errors), so a sheet of only-broken rows is not "empty file".
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["А", None], ["Б", "31.12.2025"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal
        assert plan.planned_rows == []
        assert [i.row_number for i in plan.skipped_rows] == [2, 3]


# ── 2.1 — sheet mapping and warnings ──────────────────────────────────────

class TestSheetMappingAndWarnings:
    async def test_partial_file_imports_only_known_sheets_with_warning(self, tmp_path):
        # Spec scenario: only «Персонажи» (+ unknown sheet) — no error.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["Иван", date(2001, 1, 1)]])
        _sheet(wb, "Заметки", ["Что-то"], [["..."], ["ещё"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal
        assert [row.entity_type for row in plan.planned_rows] == ["character"]
        assert any("Заметки" in w for w in plan.warnings)

    async def test_sheet_names_are_case_insensitive(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "события", CHAR_HEADERS, [["Е", date(2001, 1, 1)]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal and not plan.warnings
        assert plan.lookup_row("event", "Е") is not None

    async def test_unknown_extra_columns_ignored_silently(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS + ["Заметка"],
               [["Иван", date(2001, 1, 1), "неImportable"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal and plan.warnings == [] and plan.skipped_rows == []
        row = plan.planned_rows[0]
        assert "Заметка" not in row.fields and "note" not in row.fields

    async def test_english_legacy_headers_on_named_sheets(self, tmp_path):
        wb = _new_workbook()
        # «События» has no «Изображение» column → the legacy image header is
        # simply unknown there, everything else reads normally.
        _sheet(wb, "События", ["name", "START_DATE", "characteristics", "image"],
               [["Battle", "2001-02-03", "Big", "x.png"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal and not plan.skipped_rows
        row = plan.planned_rows[0]
        assert row.fields["name"] == "Battle"
        assert row.fields["start_date"] == MonthDay(2001, 2, 3)
        assert row.fields["characteristics"] == "Big"
        assert "image" not in row.fields

    async def test_special_and_legacy_columns_stored_by_registry_key(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События",
               ["Имя", "Дата начала", "Дата конца", "Характеристики", "Предыстория", "Рейтинг", "Тип"],
               [["Бал", "1815-01-10", "1816-06-01", "Х", "Б", "3", "Дуэль"]])
        _sheet(wb, "Персонажи",
               ["Имя", "Дата начала", "Изображение", "personality", "tasks", "music_url", "Тип"],
               [["Иван", date(2001, 1, 1), "img.png", "Смелый", "Задания", "http://m", "неНужен"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal and not plan.skipped_rows
        event = plan.lookup_row("event", "Бал")
        assert event.fields["event_type"] == "Дуэль"
        assert event.fields["rating"] == 3
        assert event.fields["start_date"] == MonthDay(1815, 1, 10)
        assert event.fields["end_date"] == MonthDay(1816, 6, 1)
        character = plan.lookup_row("character", "Иван")
        assert character.fields["image"] == "img.png"
        assert character.fields["personality"] == "Смелый"
        assert character.fields["tasks"] == "Задания"
        assert character.fields["music_url"] == "http://m"
        # «Тип» is not a Персонажи column → ignored.
        assert "event_type" not in character.fields

    async def test_link_cell_segments_normalized(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS,
               [["Е", "2001-01-01", " Иван ; ; Мария "]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.skipped_rows
        targets = [link.name for link in plan.lookup_row("event", "Е").links["character"]]
        assert targets == ["Иван", "Мария"]


# ── 2.2 — in-file merge and planned skips ─────────────────────────────────

class TestMerge:
    async def test_same_name_rows_later_non_empty_overrides(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи",
               ["Имя", "Дата начала", "Дата конца", "Характеристики", "Предыстория", "Рейтинг"],
               [
                   ["Иван", "2001-01-01", "2005-01-01", "А1", "БС1", 2],
                   ["Иван", "2002-02-02", None, "А2", None, None],
               ])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.planned_rows and not plan.skipped_rows
        merged = plan.lookup_row("character", "Иван")
        assert merged.fields["start_date"] == MonthDay(2002, 2, 2)      # later wins
        assert merged.fields["characteristics"] == "А2"             # later wins
        assert merged.fields["backstory"] == "БС1"                  # empty ≠ erase
        assert merged.fields["end_date"] == MonthDay(2005, 1, 1)        # empty ≠ erase
        assert merged.fields["rating"] == 2                         # empty ≠ erase
        assert merged.first_row_number == 2

    async def test_merge_keeps_name_case_of_the_later_row(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["Иван", "2001-01-01"], ["иван", "2002-02-02"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert len(plan.planned_rows) == 1
        merged = plan.planned_rows[0]
        assert merged.key == ("character", "иван")
        assert merged.name == "иван"          # last non-empty override
        assert merged.first_row_number == 2   # both rows merged into one

    async def test_links_accumulate_across_merged_rows(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS,
               [
                   ["Бал", "1815-01-10", "Иван"],
                   ["бал", "1815-02-10", "Мария; Иван"],
               ])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        merged = plan.lookup_row("event", "Бал")
        refs = merged.links["character"]
        assert [r.name for r in refs] == ["Иван", "Мария"]  # accumulated, first-seen order
        assert [r.source_row_number for r in refs] == [2, 3]
        assert merged.fields["start_date"] == MonthDay(1815, 2, 10)

    async def test_same_target_twice_in_one_cell_deduped(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS,
               [["Бал", "1815-01-10", "Иван;иван; Иван"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert len(plan.lookup_row("event", "Бал").links["character"]) == 1

    async def test_unparsable_end_date_ignored_not_skipped(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", ["Имя", "Дата начала", "Дата конца"],
               [["Е", "2001-01-01", "31.12.2025"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.skipped_rows
        assert "end_date" not in plan.lookup_row("event", "Е").fields

    async def test_blank_padding_rows_ignored(self, tmp_path):
        # openpyxl reads dimension-padded rows as all-empty tuples: silently
        # skipped, no «пустое имя» noise, row numbering keeps counting.
        wb = _new_workbook()
        ws = wb.create_sheet("Персонажи")
        ws.append(CHAR_HEADERS)
        ws.append(["П1", "2001-01-01"])
        ws.append([None, None])
        ws.append([None, None])
        ws.append(["П2", "2001-01-02"])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.skipped_rows == []
        assert [r.name for r in plan.planned_rows] == ["П1", "П2"]
        assert plan.lookup_row("character", "П2").first_row_number == 5


class TestPlannedSkips:
    async def test_empty_and_whitespace_names_skipped_with_position(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS,
               [[None, "2001-01-01"], ["   ", "2001-01-01"], [
                   "Иван", "2001-01-02"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert [i.row_number for i in plan.skipped_rows] == [2, 3]
        assert all(i.sheet == "Персонажи" for i in plan.skipped_rows)
        assert all("пустое имя" in i.reason for i in plan.skipped_rows)
        assert [r.name for r in plan.planned_rows] == ["Иван"]

    async def test_empty_start_date_skipped(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["БезДаты", None], ["Иван", "2001-01-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        issue, = plan.skipped_rows
        assert (issue.sheet, issue.row_number) == ("Персонажи", 2)
        assert "дата начала" in issue.reason and "пусто" in issue.reason
        assert [r.name for r in plan.planned_rows] == ["Иван"]

    async def test_unparsable_start_date_skipped(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["Иван", "31.12.2025"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        issue, = plan.skipped_rows
        assert "31.12.2025" in issue.reason and "не является датой" in issue.reason

    async def test_native_and_text_dates_parse(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "События", ["Имя", "Дата начала", "Дата конца"],
               [["Е1", date(1200, 3, 4), date(1200, 4, 5)], ["Е2", "2001-02-03", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.skipped_rows
        assert plan.lookup_row("event", "Е1").fields["start_date"] == MonthDay(1200, 3, 4)
        assert plan.lookup_row("event", "Е2").fields["start_date"] == MonthDay(2001, 2, 3)

    async def test_bc_dates_plan_with_era_flags(self, tmp_path):
        # add-era-aware-dates 5.1/5.2: both text BC forms land as (date, True)
        # next to their date; native cells keep era «н.э.».
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Имя", "Дата начала", "Дата конца"],
               [
                   ["П1", "5 марта 44 г. до н.э.", "-0001-01-01"],
                   ["П2", "-0044-03-05", None],
                   ["П3", date(44, 3, 5), None],
               ])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.skipped_rows
        p1 = plan.lookup_row("character", "П1").fields
        assert (p1["start_date"], p1["start_bc"]) == (MonthDay(44, 3, 5), True)
        assert (p1["end_date"], p1["end_bc"]) == (MonthDay(1, 1, 1), True)
        p2 = plan.lookup_row("character", "П2").fields
        assert (p2["start_date"], p2["start_bc"]) == (MonthDay(44, 3, 5), True)
        assert "end_date" not in p2 and "end_bc" not in p2
        p3 = plan.lookup_row("character", "П3").fields
        assert (p3["start_date"], p3["start_bc"]) == (MonthDay(44, 3, 5), False)

    async def test_month_less_bc_text_is_a_row_problem(self, tmp_path):
        # Месяц обязателен (spec «Колонки листа»): «44 г. до н.э.» остаётся
        # неразбираемым вводом → проблема строки, как и любая битая дата.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Имя", "Дата начала"], [["П4", "44 г. до н.э."]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        issue, = plan.skipped_rows
        assert (issue.sheet, issue.row_number) == ("Персонажи", 2)
        assert "44 г. до н.э." in issue.reason and "не является датой" in issue.reason
        assert plan.planned_rows == []

    async def test_bc_end_date_pair_survives_merge(self, tmp_path):
        # Дата конца до н.э. из ранней строки сохраняется при слиянии, а её
        # битая поздняя запись ничего не затирает (пары дата+эра атомарны).
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Имя", "Дата начала", "Дата конца"],
               [["Иван", "-0100-01-01", "-0050-12-31"],
                ["Иван", "-0099-01-01", "31.12.2025"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.skipped_rows
        merged = plan.lookup_row("character", "Иван").fields
        assert (merged["start_date"], merged["start_bc"]) == (MonthDay(99, 1, 1), True)
        assert (merged["end_date"], merged["end_bc"]) == (MonthDay(50, 12, 31), True)

    async def test_rating_out_of_range_skips_row(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS + ["Рейтинг"],
               [["Иван", "2001-01-01", 9], ["Мария", "2001-01-01", 3]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        issue, = plan.skipped_rows
        assert (issue.sheet, issue.row_number) == ("Персонажи", 2)
        assert "вне диапазона" in issue.reason
        assert plan.lookup_row("character", "Мария").fields["rating"] == 3

    async def test_rating_garbage_skips_row(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS + ["Рейтинг"], [["Иван", "2001-01-01", "abc"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        issue, = plan.skipped_rows
        assert "не целое" in issue.reason

    async def test_rating_column_exists_on_every_sheet_including_events(self, tmp_path):
        # The rating VALUE of an event row is kept in the plan fields here;
        # the events entity has no rating field, so the apply pass (group 3)
        # silently ignores this key for events — analysis must not crash.
        wb = _new_workbook()
        _sheet(wb, "События", CHAR_HEADERS + ["Рейтинг"], [["Е", "2001-01-01", 4]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.planned_rows[0].fields["rating"] == 4


# ── 2.3 — name index against an in-memory DB ──────────────────────────────

async def _seed_db(async_session):
    """Two «Иван» characters (case-varied duplicate), one «Уникал», one item."""
    ivan_a = CharacterModel(name="Иван", start_date=date(1900, 1, 1))
    ivan_b = CharacterModel(name="иван", start_date=date(1901, 1, 1))
    unique_char = CharacterModel(name="Уникал", start_date=date(1902, 1, 1))
    item = ItemModel(name="Фонарь", start_date=date(1905, 1, 1))
    async_session.add_all([ivan_a, ivan_b, unique_char, item])
    await async_session.flush()
    return unique_char, item


class TestNameIndex:
    async def test_unique_db_match_becomes_update_plan_with_nonempty_only(
        self, tmp_path, async_session
    ):
        unique_char, _ = await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Имя", "Дата начала", "Характеристики", "Предыстория"],
               [["Уникал", "2001-01-01", "новое", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        row = plan.lookup_row("character", "Уникал")
        assert row.is_update and row.existing_id == unique_char.id
        # Update touches only non-empty cells (spec «Пустая ячейка не затирает»).
        assert row.fields["characteristics"] == "новое"
        assert "backstory" not in row.fields

    async def test_unique_match_is_case_insensitive(self, tmp_path, async_session):
        unique_char, _ = await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["уникал", "2001-01-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        row = plan.lookup_row("character", "уникал")
        assert row.existing_id == unique_char.id

    async def test_ambiguous_db_name_row_becomes_create_plan(self, tmp_path, async_session):
        await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["иван", "2001-01-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        row = plan.lookup_row("character", "иван")
        assert not row.is_update and row.existing_id is None
        # The GROUP BY lower(name) HAVING COUNT(*)>1 result is visible on the plan.
        assert plan.ambiguous_names["character"] == {"иван"}

    async def test_ambiguous_link_reference_skips_referring_row(self, tmp_path, async_session):
        await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "иван"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        assert plan.planned_rows == []
        issue, = plan.skipped_rows
        assert (issue.sheet, issue.row_number) == ("События", 2)
        assert "иван" in issue.reason and "разрешения не имеет" in issue.reason
        assert plan.ghosts == {}

    async def test_file_row_beats_ambiguous_db_name(self, tmp_path, async_session):
        await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "иван"]])
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["Иван", "1800-01-01"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        event = plan.lookup_row("event", "Бал")
        assert not event.skipped           # file priority → no planned skip
        link = event.links["character"][0]
        assert link.resolution == LINK_TO_FILE
        assert link.target_key == plan.lookup_row("character", "Иван").key
        # The character itself still creates anew (DB name ambiguous).
        assert plan.lookup_row("character", "Иван").existing_id is None
        assert plan.ghosts == {}

    async def test_unique_db_link_target_resolves_to_id(self, tmp_path, async_session):
        _, item = await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "События", ["Имя", "Дата начала", "Связь предметами"],
               [["Бал", "1815-01-10", "фонарь"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        link = plan.lookup_row("event", "Бал").links["item"][0]
        assert link.resolution == LINK_TO_DB
        assert link.db_id == item.id

    async def test_unknown_link_target_becomes_ghost_with_referring_dates(
        self, tmp_path, async_session
    ):
        await _seed_db(async_session)
        wb = _new_workbook()
        # Spec scenario: «Бал» 1820-05-01 (no end) → ghost item 1820-05-01–1820-05-01.
        _sheet(wb, "События", ["Имя", "Дата начала", "Связь предметами", "Дата конца"],
               [["Бал", "1820-05-01", "Амулет", None]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        link = plan.lookup_row("event", "Бал").links["item"][0]
        assert link.resolution == LINK_TO_GHOST and link.db_id is None
        ghost = plan.ghosts[("item", "амулет")]
        assert (ghost.min_start, ghost.max_end) == (MonthDay(1820, 5, 1), MonthDay(1820, 5, 1))
        assert ghost.name == "Амулет"
        assert ghost.referenced_by == [("События", 2)]

    async def test_ghost_dates_span_min_start_max_end_over_references(
        self, tmp_path, async_session
    ):
        await _seed_db(async_session)
        wb = _new_workbook()
        # Spec scenario «Противоречивые даты»: 1820-05-01 + 1815-01-10 → 1815-01-10..1820-05-01.
        _sheet(wb, "События", ["Имя", "Дата начала", "Дата конца", "Связь предметами"],
               [
                   ["Бал", "1820-05-01", None, "Амулет"],
                   ["Охота", "1815-01-10", "1816-06-01", "Амулет"],
               ])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        ghost = plan.ghosts[("item", "амулет")]
        assert ghost.min_start == MonthDay(1815, 1, 10)
        assert ghost.max_end == MonthDay(1820, 5, 1)
        assert ghost.referenced_by == [("События", 2), ("События", 3)]

    async def test_ghost_bounds_compare_through_the_era_key(self, tmp_path, async_session):
        # add-era-aware-dates 5.2: min/max считаются ключом эпохи. Наивное
        # сравнение date-объектов взяло бы 100 г. до н.э. раньше 500 г. до н.э.
        # и потеряло бы эру конца; через ключ min — 500 г. до н.э.,
        # max — любой конец н.э.
        await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "События", ["Имя", "Дата начала", "Дата конца", "Связь предметами"],
               [
                   ["Позже в до н.э.", "-0100-01-01", None, "Амулет"],
                   ["Раньше в до н.э.", "-0500-06-01", "-0499-12-31", "Амулет"],
                   ["Наша эра", "2026-08-01", None, "Амулет"],
               ])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        ghost = plan.ghosts[("item", "амулет")]
        assert (ghost.min_start, ghost.min_start_bc) == (MonthDay(500, 6, 1), True)
        assert (ghost.max_end, ghost.max_end_bc) == (MonthDay(2026, 8, 1), False)

    async def test_link_lookup_is_type_scoped_against_db(self, tmp_path, async_session):
        _, item = await _seed_db(async_session)  # «Фонарь» is an ITEM in DB
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "Фонарь"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        link = plan.lookup_row("event", "Бал").links["character"][0]
        # Same name exists in DB but of another type → not a match → ghost.
        assert link.resolution == LINK_TO_GHOST
        assert link.target_key in plan.ghosts

    async def test_link_reference_against_empty_db_becomes_ghost_not_fatal(
        self, tmp_path, async_session
    ):
        # The DB was consulted, found nothing → a legitimate ghost, even
        # though the name index stays empty (built, just empty).
        wb = _new_workbook()
        _sheet(wb, "События", ["Имя", "Дата начала", "Связь предметами"],
               [["Бал", "1820-05-01", "Фонарь"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        assert not plan.has_fatal
        assert plan.name_index_built and plan.name_index == {}
        assert plan.ghosts[("item", "фонарь")].min_start == MonthDay(1820, 5, 1)

    async def test_cascade_when_file_target_row_gets_skipped(self, tmp_path, async_session):
        # «Глеб» (character) links «Цех», which is ambiguous in DB and absent
        # from the file → Глеб's row is skipped; the event that links to Глеб
        # must then re-resolve to a ghost, not to the vanished file row
        # (fixpoint resolution over the growing skip set).
        await _seed_db(async_session)
        org_a = OrganizationModel(name="Цех", start_date=date(1850, 1, 1))
        org_b = OrganizationModel(name="цех", start_date=date(1851, 1, 1))
        async_session.add_all([org_a, org_b])
        await async_session.flush()
        wb = _new_workbook()
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "Глеб"]])
        _sheet(wb, "Персонажи", ["Имя", "Дата начала", "Связь организациями"],
               [["Глеб", "1800-01-01", "цех"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))

        gleb = plan.lookup_row("character", "Глеб")
        assert gleb.skipped                            # its own ambiguous reference
        event = plan.lookup_row("event", "Бал")
        assert not event.skipped
        link = event.links["character"][0]
        assert link.resolution == LINK_TO_GHOST
        # No «Глеб» in DB → the ghost takes the event's dates.
        assert plan.ghosts[("character", "глеб")].min_start == MonthDay(1815, 1, 10)
        assert plan.ambiguous_names["organization"] == {"цех"}
        reasons = [(i.sheet, i.row_number) for i in plan.skipped_rows]
        assert reasons == [("Персонажи", 2)]
        # «Глеб»-ссылка события уцелевшего, ссылки пропущенной строки — вне плана.
        assert [link.name for link in plan.iter_links()] == ["Глеб"]

    async def test_analyze_writes_nothing_to_the_db(self, tmp_path, async_session):
        await _seed_db(async_session)
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["Новый", "2001-01-01"]])
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "иван"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb), GameSessionUoW(async_session))
        assert not plan.has_fatal  # some rows skipped, the DB itself untouched
        chars = (await async_session.execute(
            select(func.count()).select_from(CharacterModel)
        )).scalar()
        items = (await async_session.execute(
            select(func.count()).select_from(ItemModel)
        )).scalar()
        assert chars == 3 and items == 1


# ── 2.3 — analysis without a DB session ───────────────────────────────────

class TestAnalysisWithoutSession:
    async def test_file_and_ghost_resolution_works_without_db(self, tmp_path):
        wb = _new_workbook()
        _sheet(wb, "Персонажи", CHAR_HEADERS, [["П1", "2001-01-01"]])
        _sheet(wb, "События", EVENT_HEADERS, [["Бал", "1815-01-10", "П1; Призрак"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert not plan.has_fatal
        links = plan.lookup_row("event", "Бал").links["character"]
        assert [link.resolution for link in links] == [LINK_TO_FILE, LINK_TO_GHOST]
        assert plan.ghosts[("character", "призрак")].min_start == MonthDay(1815, 1, 10)

    async def test_unverifiable_reference_without_db_is_fatal(self, tmp_path):
        # Link column of a type the file never defines: without the DB the
        # reference cannot be classified (DB match vs ghost) → fatal, so no
        # silently wrong auto-created entity.
        wb = _new_workbook()
        _sheet(wb, "Персонажи", ["Имя", "Дата начала", "Связь предметами"],
               [["Иван", "2001-01-01", "Фонарь"]])
        plan = await _svc().analyze_file(_save(tmp_path, wb))
        assert plan.has_fatal
        assert "индекс имён" in plan.fatal_errors[0]
