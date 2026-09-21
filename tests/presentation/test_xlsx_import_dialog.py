"""Characterization tests for the unified XlsxImportDialog island (task 4.2).

The dialog is the render half of the VM state machine (task 4.1): pre-analysis
problem list (sheet/row/reason) with the primary button blocked on fatals,
progress, the final report panel, the registry-generated hint, the «Скачать
шаблон» button and the honest `.xlsx`-only file filter.
"""
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QFileDialog, QMessageBox

from app.application.services import xlsx_schema
from app.application.services.xlsx_import_service import (
    DateShiftRow,
    ImportPlan,
    ImportReport,
    RowIssue,
)
from app.application.services.xlsx_template import build_template_workbook
from app.domain.game_calendar import (
    CalendarSpec,
    CustomCalendar,
    IntercalarySpec,
    MonthSpec,
    StandardCalendar,
    current_calendar,
    reset_current_calendar,
    set_current_calendar,
)
from app.presentation.views import xlsx_import_dialog
from app.presentation.views.xlsx_import_dialog import (
    PRIMARY_TEXT_ANALYZE,
    PRIMARY_TEXT_CLOSE,
    PRIMARY_TEXT_IMPORT,
    TEMPLATE_FILE_NAME,
    XLSX_FILTER,
    XlsxImportDialog,
    build_format_text,
    date_formats_hint,
    save_template_as,
)
from tests.presentation.qml_helpers import find_item, island_row_texts, walk_items


def _cell_table(wb: Workbook, title: str) -> list[list[object]]:
    """Cell values of one sheet with datetimes normalized to dates."""
    return [
        [value.date() if isinstance(value, datetime) else value for value in raw]
        for raw in wb[title].iter_rows(values_only=True)
    ]


#: The «Даты» block of the preset hint, word-for-word (spec «Подсказка о
#: колонках в диалоге» continuity, design D9) — the tests below forbid drift.
PRESET_DATE_BLOCK = (
    "Даты: нативная ячейка Excel или текст YYYY-MM-DD — это наша эра;\n"
    "до н.э. — текст «5 марта 44 г. до н.э.» или знаковое ISO -0044-03-05."
)


def _hint_calendar(*, first_month_days: int = 30, intercalary: bool = True) -> CustomCalendar:
    """Custom calendar for the hint tests: «Зимостой» opens the year, the
    «Медожор» intercalary day rides after the second month when declared."""
    return CustomCalendar(CalendarSpec(
        months=(
            MonthSpec("Зимостой", first_month_days),
            MonthSpec("Вьюжень", 9),
            MonthSpec("Травень", 12),
        ),
        week_names=("рысь", "волк", "лиса", "лось", "барс", "соня", "зверь", "ёж"),
        intercalary=(IntercalarySpec("Медожор", after_month=2),) if intercalary else (),
    ))


def plan_with(*, fatal=(), rows=(), warnings=()) -> ImportPlan:
    plan = ImportPlan(path="/tmp/file.xlsx")
    plan.fatal_errors = list(fatal)
    plan.skipped_rows = list(rows)
    plan.warnings = list(warnings)
    return plan


@pytest.fixture
def dlg(qtbot):
    d = XlsxImportDialog()
    qtbot.addWidget(d)
    return d


def _texts(widget, object_name: str) -> list[str]:
    return [i.property("text") for i in walk_items(widget.rootObject())
            if i.objectName() == object_name]


class TestHintFromRegistry:
    def test_hint_covers_every_sheet_and_column(self):
        hint = build_format_text()
        for sheet in xlsx_schema.all_sheets():
            assert f"Лист «{sheet.sheet_name}»" in hint
            for column in xlsx_schema.all_headers(sheet):
                assert column.label in hint
        # Continuity of the old hint wording for the shared columns; the
        # date wording additionally names the BC forms (era-aware dates).
        assert "Дата начала" in hint and "дата Excel" in hint
        assert "5 марта 44 г. до н.э." in hint
        assert "-0044-03-05" in hint
        assert "Ссылка на музыкальную тему" in hint
        assert "PNG, JPG, BMP, GIF, WebP" in hint

    def test_hint_shows_link_syntax_and_aliases(self):
        hint = build_format_text()
        assert "Связь персонажами" in hint
        assert "Связь событиями" in hint
        assert "имена через «;»" in hint
        assert "алиасы: name" in hint  # старый англ. заголовок упомянут

    def test_dialog_hint_matches_registry_generation(self, dlg):
        assert dlg.format_text.toPlainText() == build_format_text()
        assert dlg.format_text.isReadOnly()


class TestDateFormatsHint:
    """C5 task 5.1 — the «Даты» helper: the preset answers the old two lines
    word-for-word; a custom calendar keeps them and adds the game wording with
    a valid example of ITS calendar (day = min(3, длина первого месяца)),
    naming a declared intercalary day only when the calendar declares one."""

    def test_preset_answer_is_the_two_old_lines_word_for_word(self):
        assert date_formats_hint(StandardCalendar()) == PRESET_DATE_BLOCK

    def test_custom_keeps_old_lines_and_adds_game_example(self):
        calendar = _hint_calendar()
        set_current_calendar(calendar)
        try:
            lines = date_formats_hint(calendar).splitlines()
        finally:
            reset_current_calendar()
        assert lines[:2] == PRESET_DATE_BLOCK.splitlines()  # прежние строки дословно
        assert "«03 Зимостой 44 г. до н.э.»" in lines[2]  # валидный пример активного календаря
        assert "«Медожор 44 г. до н.э.»" in lines[2]  # имя объявленного вставного дня

    def test_custom_day_clamps_to_short_first_month(self):
        # min(3, длина первого месяца): a two-day «Зимостой» cannot advertise a third day.
        calendar = _hint_calendar(first_month_days=2, intercalary=False)
        set_current_calendar(calendar)
        try:
            hint = date_formats_hint(calendar)
        finally:
            reset_current_calendar()
        assert "«02 Зимостой 44 г. до н.э.»" in hint

    def test_custom_without_intercalary_omits_intercalary_day(self):
        calendar = _hint_calendar(intercalary=False)
        set_current_calendar(calendar)
        try:
            hint = date_formats_hint(calendar)
        finally:
            reset_current_calendar()
        assert "«03 Зимостой 44 г. до н.э.»" in hint
        assert "вставной" not in hint and "Медожор" not in hint


class TestHintDateBlockIsCalendarAware:
    """C5 task 5.2 — build_format_text substitutes the calendar-aware block and
    the dialog overrides the date-column descriptions: preset stays word-for-word
    (spec «Пользователь видит спецификацию» + continuity), a custom game sees the
    «Зимостой» example alongside the ISO forms (spec «Подсказка кастомной игры
    показывает игровые формы»)."""

    def test_preset_hint_is_word_for_word(self):
        hint = build_format_text()
        assert (
            "старые английские заголовки читаются как алиасы.\n\n"
            + PRESET_DATE_BLOCK + "\nСвязи в ячейке"
        ) in hint
        assert (
            "| да  | Начало: YYYY-MM-DD, дата Excel (наша эра), "
            "«5 марта 44 г. до н.э.» или -0044-03-05 (алиасы: start_date)"
        ) in hint
        assert "| нет | Конец: форматы те же, что у даты начала" in hint

    def test_custom_block_and_start_column_get_the_game_example(self):
        calendar = _hint_calendar()
        set_current_calendar(calendar)
        try:
            hint = build_format_text()
        finally:
            reset_current_calendar()
        # Блок «Даты»: игровая форма наряду с ISO-формами прежних строк.
        assert "«03 Зимостой 44 г. до н.э.»" in hint
        assert "YYYY-MM-DD" in hint and "-0044-03-05" in hint
        # Подпись «Дата начала» описывает тот же набор форм, что парсер ветки.
        assert (
            "| да  | Начало: YYYY-MM-DD, дата Excel (наша эра), "
            "«5 марта 44 г. до н.э.» или -0044-03-05, игровая форма "
            "«03 Зимостой 44 г. до н.э.» (вставной день — «Медожор 44 г. до н.э.»)"
            " (алиасы: start_date)"
        ) in hint
        # «Дата конца» описана относительно — формулировка держится в обеих ветках.
        assert "| нет | Конец: форматы те же, что у даты начала" in hint

    def test_dialog_renders_the_custom_calendar_hint(self, qtbot):
        # Диалог перенимает календарь-зависимый текст при конструировании.
        calendar = _hint_calendar()
        set_current_calendar(calendar)
        try:
            d = XlsxImportDialog()
            qtbot.addWidget(d)
            text = d.format_text.toPlainText()
        finally:
            reset_current_calendar()
        assert "«03 Зимостой 44 г. до н.э.»" in text
        assert "игровая форма" in text


class TestConstruction:
    def test_title_and_initial_state(self, dlg):
        assert dlg.windowTitle() == "Импорт из .xlsx"
        assert dlg.get_path() == ""
        assert dlg.path_edit.text() == ""
        assert dlg.progress_bar.value() == 0
        assert dlg.import_btn.isEnabled() is False  # no path in idle
        assert dlg.import_btn.text() == PRIMARY_TEXT_ANALYZE
        assert dlg.vm.state == "idle"
        assert not dlg.vm.hasFatal and dlg.vm.report is None

    def test_island_object_names(self, dlg):
        root = dlg.quick.rootObject()
        names = {root.objectName()} | {i.objectName() for i in walk_items(root)}
        for name in (
            "formatArea", "pathField", "browseButton", "downloadButton",
            "progressBar", "importButton", "cancelButton", "issueList",
            "reportArea",
        ):
            assert name in names
        assert not hasattr(dlg, "entity_type")
        assert root.property("defaultButton").objectName() == "importButton"


class TestBrowseFilter:
    def test_filter_is_xlsx_only(self, dlg, mocker):
        spy = mocker.patch.object(QFileDialog, "getOpenFileName", return_value=("", ""))
        dlg._on_browse()
        selected_filter = spy.call_args.args[3]
        assert selected_filter == XLSX_FILTER
        assert "*.xlsx" in selected_filter
        assert "*.xls " not in selected_filter  # .xls честно исключён
        assert "Все файлы" not in selected_filter

    def test_browse_sets_path(self, dlg, mocker, tmp_path):
        p = tmp_path / "file.xlsx"
        p.write_bytes(b"x")
        mocker.patch.object(QFileDialog, "getOpenFileName", return_value=(str(p), ""))
        dlg.vm.requestBrowse()
        assert dlg.path_edit.text() == str(p)
        assert dlg.get_path() == str(p)

    def test_browse_cancellation_keeps_path_empty(self, dlg, mocker):
        mocker.patch.object(QFileDialog, "getOpenFileName", return_value=("", ""))
        dlg.vm.requestBrowse()
        assert dlg.path_edit.text() == ""
        assert dlg.get_path() == ""


class TestAnalyzeFlow:
    def test_primary_click_requests_analysis(self, dlg):
        dlg.path_edit.setText("/tmp/batch.xlsx")
        with qtbot_signal(dlg, "analyze_requested") as received:
            dlg.import_btn.click()
        assert received == ["/tmp/batch.xlsx"]
        assert dlg.vm.state == "analyzing"
        assert dlg.import_btn.isEnabled() is False

    def test_fatal_plan_blocks_button_and_lists_reason(self, dlg):
        dlg.path_edit.setText("/tmp/broken.xlsx")
        dlg.vm.on_analyzed(plan_with(fatal=["Лист «Персонажи»: нет строки заголовков."]))
        assert dlg.import_btn.isEnabled() is False
        assert dlg.import_btn.text() == PRIMARY_TEXT_IMPORT  # состояние problems…
        dlg.import_btn.click()  # …но клик заблокирован стейт-машиной
        assert dlg.vm.state == "problems"

        dlg.quick.grab()
        assert find_item(dlg.quick, "fatalTitle").property("visible") is True
        reasons = island_row_texts(dlg.quick, "issueRow", "issueReasonText")
        assert reasons == ["Фатально: Лист «Персонажи»: нет строки заголовков."]
        sheets = island_row_texts(dlg.quick, "issueRow", "issueSheetText")
        assert sheets == ["файл"]

    def test_row_issues_shown_sheet_row_reason(self, dlg):
        dlg.path_edit.setText("/tmp/partial.xlsx")
        dlg.vm.on_analyzed(plan_with(rows=[
            RowIssue("Персонажи", 5, "пустое имя"),
            RowIssue("События", 9, "рейтинг 7 вне диапазона 1..5"),
        ]))
        assert dlg.import_btn.isEnabled() is True
        dlg.quick.grab()
        assert island_row_texts(dlg.quick, "issueRow", "issueSheetText") == [
            "лист «Персонажи»", "лист «События»",
        ]
        assert island_row_texts(dlg.quick, "issueRow", "issueRowNumberText") == [
            "строка 5", "строка 9",
        ]
        assert island_row_texts(dlg.quick, "issueRow", "issueReasonText") == [
            "пустое имя", "рейтинг 7 вне диапазона 1..5",
        ]

    def test_confirm_from_problems_emits_signal_and_shows_progress(self, dlg):
        dlg.path_edit.setText("/tmp/partial.xlsx")
        dlg.vm.on_analyzed(plan_with(rows=[RowIssue("Персонажи", 5, "пустое имя")]))
        assert dlg.import_btn.text() == PRIMARY_TEXT_IMPORT
        with qtbot_signal(dlg, "confirm_import") as received:
            dlg.import_btn.click()
        assert received == [()]
        assert dlg.vm.state == "importing"
        assert find_item(dlg.quick, "progressBar").property("visible") is True


class TestProgress:
    def test_set_progress_partial_and_complete(self, dlg):
        dlg.path_edit.setText("/tmp/a.xlsx")
        dlg.vm.on_analyzed(plan_with())
        dlg.vm.requestConfirmImport()
        dlg.set_progress(1, 4)
        assert dlg.progress_bar.value() == 25
        dlg.set_progress(3, 4)
        assert dlg.progress_bar.value() == 75
        dlg.set_progress(4, 4)
        assert dlg.progress_bar.value() == 100

    def test_set_progress_zero_total_resets(self, dlg):
        dlg.set_progress(2, 2)
        dlg.set_progress(0, 0)
        assert dlg.progress_bar.value() == 0

    def test_progress_visibility_mirrors_view_model(self, dlg):
        assert dlg.progress_bar.isVisible() is False
        dlg.path_edit.setText("/tmp/a.xlsx")
        dlg.vm.on_analyzed(plan_with())
        dlg.vm.requestConfirmImport()  # importing → полоса видима
        assert dlg.progress_bar.isVisible() is True


class TestPlanOwnership:
    def test_set_plan_adopts_the_analyzed_plan(self, dlg):
        plan = plan_with()
        dlg.set_plan(plan)
        assert dlg.plan is plan


class TestReportPanel:
    @pytest.fixture
    def reported(self, dlg):
        dlg.path_edit.setText("/tmp/a.xlsx")
        report = ImportReport(
            created=8,
            updated=2,
            links=5,
            skipped=[RowIssue("Персонажи", 3, "пустое имя")],
            decisions=["Автосоздание: тип события «Дуэль»"],
            warnings=["Неизвестный лист «Заметки»"],
        )
        dlg.vm.on_analyzed(plan_with(rows=report.skipped))
        dlg.vm.requestConfirmImport()
        dlg.vm.on_report(report)
        return dlg, report

    def test_state_and_button(self, reported):
        d, _ = reported
        assert d.vm.state == "done"
        assert d.import_btn.isEnabled() is True
        assert d.import_btn.text() == PRIMARY_TEXT_CLOSE

    def test_panel_numbers_and_lists(self, reported):
        d, report = reported
        d.quick.grab()
        assert find_item(d.quick, "reportArea").property("visible") is True
        assert find_item(d.quick, "reportCreatedText").property("text") == "Создано: 8"
        assert find_item(d.quick, "reportUpdatedText").property("text") == "Обновлено: 2"
        assert find_item(d.quick, "reportLinksText").property("text") == "Связей: 5"
        assert "Пропущено строк: 1" in _texts(d.quick, "reportSkippedTitle")
        assert _texts(d.quick, "reportSkippedRowText") == [
            "лист «Персонажи», строка 3: пустое имя"
        ]
        assert _texts(d.quick, "reportDecisionText") == [
            "Решение: Автосоздание: тип события «Дуэль»"
        ]
        assert _texts(d.quick, "reportWarningText") == [
            "Предупреждение: Неизвестный лист «Заметки»"
        ]

    def test_close_via_primary_click(self, reported, qtbot):
        d, _ = reported
        with qtbot.waitSignal(d.finished, timeout=2000):
            d.import_btn.click()

    def test_date_shifts_section_hidden_when_empty(self, reported):
        # C3a task 4.2: no actual transfers (standard game) — no section.
        d, _ = reported
        d.quick.grab()
        title = find_item(d.quick, "reportDateShiftTitle")
        assert title is not None and title.property("visible") is False
        assert _texts(d.quick, "reportDateShiftText") == []

    def test_date_shifts_section_lists_transfers(self, dlg):
        dlg.path_edit.setText("/tmp/a.xlsx")
        report = ImportReport(
            created=1,
            date_shifts=[
                DateShiftRow("Персонажи", 2, "Дата начала", "2026-08-31", "2026-08-30"),
            ],
        )
        dlg.vm.on_analyzed(plan_with())
        dlg.vm.requestConfirmImport()
        dlg.vm.on_report(report)
        dlg.quick.grab()
        assert _texts(dlg.quick, "reportDateShiftTitle") == ["Перенесённые даты: 1"]
        assert _texts(dlg.quick, "reportDateShiftText") == [
            "лист «Персонажи», строка 2: «Дата начала» 2026-08-31 → 2026-08-30"
        ]


class TestDownloadTemplate:
    def test_button_emits_dialog_signal(self, dlg, qtbot):
        # Seam for task 5.2: the island only emits, the save--as flow lands there.
        with qtbot.waitSignal(dlg.download_template, timeout=2000):
            dlg.vm.requestDownloadTemplate()


class TestDownloadTemplateFlow:
    """C5 task 4.1 — the save--as side of the «Скачать шаблон» seam generates
    the template (design D8): nothing is built while the user still decides,
    and what lands on disk is the workbook of the calendar active at the very
    moment of saving (no static bundle file anymore)."""

    def test_get_save_file_name_defaults_and_filter(self, mocker, tmp_path):
        out = tmp_path / "мой_шаблон.xlsx"
        spy = mocker.patch.object(
            QFileDialog, "getSaveFileName", return_value=(str(out), XLSX_FILTER),
        )
        written = save_template_as(None)
        assert spy.call_args.args[2] == TEMPLATE_FILE_NAME  # имя по умолчанию
        assert spy.call_args.args[3] == XLSX_FILTER  # честный .xlsx-фильтр
        assert written == out
        generated = build_template_workbook(current_calendar())
        downloaded = load_workbook(out, read_only=True)
        try:
            assert downloaded.sheetnames == generated.sheetnames
            for title in generated.sheetnames:
                assert _cell_table(downloaded, title) == _cell_table(generated, title)
        finally:
            downloaded.close()

    def test_template_follows_the_active_calendar(self, mocker, tmp_path):
        # The workbook answer the GAME calendar at save time (spec «Шаблон
        # сохранён» + «Шаблон кастомной игры говорит на её календаре»): a
        # custom calendar with an intercalary day puts its game wording into
        # the downloaded file — cell content no static file could pre-bake.
        out = tmp_path / "custom.xlsx"
        mocker.patch.object(QFileDialog, "getSaveFileName", return_value=(str(out), ""))
        calendar = CustomCalendar(CalendarSpec(
            months=(
                MonthSpec("Зимостой", 30), MonthSpec("Вьюжень", 9),
                MonthSpec("Травень", 12), MonthSpec("Цветень", 28),
                MonthSpec("Жневень", 20), MonthSpec("Сенокос", 33),
                MonthSpec("Гридень", 31), MonthSpec("Листопад", 7),
                MonthSpec("Хмурень", 25), MonthSpec("Студень", 44),
                MonthSpec("Крещень", 11), MonthSpec("Медовик", 21),
                MonthSpec("Чернолист", 30),
            ),
            week_names=("рысь", "волк", "лиса", "лось", "барс", "соня", "зверь", "ёж"),
            intercalary=(IntercalarySpec("Медожор", after_month=10),),
        ))
        set_current_calendar(calendar)
        try:
            written = save_template_as(None)
        finally:
            reset_current_calendar()
        assert written == out
        wb = load_workbook(out, read_only=True)
        try:
            event_dates = [row[1] for row in _cell_table(wb, "События")[1:]]
        finally:
            wb.close()
        assert "Медожор 44" in event_dates  # declared intercalary day
        assert "12 Травень 44 г. до н.э." in event_dates  # clamped game-wording BC
        assert "15 марта 44 г. до н.э." not in event_dates  # preset wording gone

    def test_cancellation_writes_nothing(self, mocker, tmp_path):
        spy = mocker.patch.object(QFileDialog, "getSaveFileName", return_value=("", ""))
        build = mocker.spy(xlsx_import_dialog, "build_template_workbook")
        assert save_template_as(None) is None
        assert spy.called
        assert list(tmp_path.iterdir()) == []
        assert not build.called  # отмена не должна порождать работу (D8)

    def test_missing_extension_gets_xlsx_appended(self, mocker, tmp_path):
        out = tmp_path / "template"  # пользователь не дописал расширение
        mocker.patch.object(QFileDialog, "getSaveFileName", return_value=(str(out), ""))
        written = save_template_as(None)
        assert written == tmp_path / "template.xlsx"
        assert written.is_file()

    def test_generation_failure_reports_to_user(self, mocker, tmp_path):
        # Broken generator (the old «resource missing» report role) — the user
        # hears about it, and no half-written file is left behind.
        out = tmp_path / "t.xlsx"
        mocker.patch.object(QFileDialog, "getSaveFileName", return_value=(str(out), ""))
        fail = mocker.patch.object(
            xlsx_import_dialog, "build_template_workbook",
            side_effect=RuntimeError("сборка упала"),
        )
        critical = mocker.patch.object(QMessageBox, "critical")
        assert save_template_as(None) is None
        assert fail.called and critical.called
        assert not out.exists()

    def test_write_failure_reports_to_user(self, mocker, tmp_path):
        # Диск под недоступным путём (прежняя ветка OSError, без copyfile).
        mocker.patch.object(
            QFileDialog, "getSaveFileName",
            return_value=(str(tmp_path / "copy" / "t.xlsx"), ""),
        )
        critical = mocker.patch.object(QMessageBox, "critical")
        assert save_template_as(None) is None
        assert critical.called


class TestPathInvalidation:
    def test_new_path_clears_problems_and_report(self, dlg):
        dlg.path_edit.setText("/tmp/broken.xlsx")
        dlg.vm.on_analyzed(plan_with(fatal=["битый"]))
        assert dlg.import_btn.isEnabled() is False
        dlg.path_edit.setText("/tmp/next.xlsx")
        assert dlg.vm.state == "idle"
        assert dlg.vm.analysisIssues == []
        assert dlg.import_btn.isEnabled() is True


class TestWiringPublishSeam:
    """The wiring feeds analyze/confirm results back through the guarded
    publish_* seam (analyze and confirm run as detached tasks; the user may
    close the dialog — «finished → deleteLater» — while they are awaiting).
    The live part lives here; the dead part deletes the C++ side for real."""

    def test_publish_analysis_adopts_plan_and_feeds_vm(self, dlg):
        plan = plan_with(rows=[RowIssue("Персонажи", 5, "пустое имя")])
        dlg.publish_analysis(plan)
        assert dlg.plan is plan
        assert dlg.vm.state == "problems"
        assert dlg.vm.analysisIssues[0]["sheet"] == "Персонажи"

    def test_publish_failure_report_transitions(self, dlg):
        dlg.publish_analysis_failure("диск отвалился")
        assert dlg.vm.state == "problems" and dlg.vm.hasFatal
        # another analysis round, then results of the apply pass
        dlg.path_edit.setText("/tmp/b.xlsx")
        dlg.publish_analysis(plan_with())
        dlg.publish_import_failure("сбой прохода 2")
        assert dlg.vm.hasFatal and "сбой" in dlg.vm.analysisIssues[0]["reason"]
        dlg.publish_report(ImportReport(created=1))
        assert dlg.vm.state == "done" and dlg.vm.report["created"] == 1

    def test_dead_dialog_publishing_and_progress_are_noops(self):
        # Mirrors the wiring teardown («finished → deleteLater») and the
        # engine.py island-release guard precedent: a deferred callback on a
        # destroyed C++ side must silently drop, not raise a RuntimeError.
        import shiboken6
        from PySide6.QtCore import QCoreApplication, QEvent
        from PySide6.QtWidgets import QApplication

        dead = XlsxImportDialog()
        dead.publish_analysis(plan_with())          # still alive: adopted
        assert dead.plan is not None
        dead.done(0)
        QApplication.processEvents()                # deferred island release
        dead.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        assert not shiboken6.isValid(dead)          # the C++ side is really gone
        plan_after = plan_with(rows=[RowIssue("События", 2, "х")])
        dead.set_progress(1, 2)                     # progress tick mid-import
        dead.publish_analysis(plan_after)           # must not resurrect/raise
        dead.publish_analysis_failure("x")
        dead.publish_report(ImportReport())
        dead.publish_import_failure("x")
        assert dead.plan is not plan_after          # publishes were dropped


# small helper: collect a named dialog signal's emits as a context manager
class qtbot_signal:
    def __init__(self, dialog, signal_name):
        self._signal = getattr(dialog, signal_name)
        self.received = []

    def __enter__(self):
        self._signal.connect(lambda *args: self.received.append(args[0] if len(args) == 1 else args))
        return self.received

    def __exit__(self, *exc):
        self._signal.disconnect()
        return False
