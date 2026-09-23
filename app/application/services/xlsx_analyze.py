"""Analyze half of the xlsx import (audit B6, design D6): workbook reading,
the merge engine, the DB name index assembly and link resolution.

Produces the pure :class:`ImportPlan` the service hands to
``xlsx_apply`` — nothing here ever writes to the database; the only DB
access is the read-only name index through the repository layer.  The code
moved byte-for-byte from the former ``xlsx_import_service`` monolith
(findings B6), only the DB queries themselves live in
``app/infrastructure/repositories/name_index.py`` now.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.application.services import xlsx_schema as schema
from app.application.services.xlsx_report_text import (
    ambiguous_link_reference,
    date_problem_caption,
)
from app.domain.date_era import era_key
from app.domain.game_calendar import (
    GameCoord,
    MonthDay,
    current_calendar,
    shift_invalid,
)
from app.infrastructure.calendar_storage import (
    encode_coord,
)
from app.infrastructure.repositories.name_index import query_lower_name_index

# Link-reference resolution kinds (design D3): resolved to a row of the same
# file, to a uniquely named DB entity, or auto-created from the reference.
LINK_TO_FILE = "file"
LINK_TO_DB = "db"
LINK_TO_GHOST = "ghost"


# ── Import plan (rework-xlsx-import, design D2) ───────────────────────────
# The plan is a pure artifact: no DB mutation happens during analysis, and
# applying it (task group 3) consumes exactly what is stored here. Fatal file
# errors block applying; row-level problems are recorded as planned skips.

@dataclass
class RowIssue:
    """One planned-skip reason for one source row (list: sheet + row number)."""

    sheet: str
    row_number: int  # 1-based Excel row number (header row = 1)
    reason: str


@dataclass(frozen=True)
class DateShiftRow:
    """One visible import transfer (piece C3a, design D7).

    A parsed number existed as a coordinate of no calendar date, so the import
    clamped it to the last valid day of the same month (or of the last existing
    one) against the *active* calendar.  The row names the source (sheet,
    1-based row number), the registry column label and the ``old``/``new``
    display strings; the era marker of BC dates rides on both.
    """

    sheet: str
    row_number: int
    field: str
    old: str
    new: str


@dataclass
class NameIndexEntry:
    """DB entities of one type sharing one lower(name) (design D3)."""

    models: list[Any] = field(default_factory=list)

    @property
    def is_unique(self) -> bool:
        return len(self.models) == 1

    @property
    def is_ambiguous(self) -> bool:
        return len(self.models) > 1


@dataclass
class PlannedLink:
    """One link-column reference collected during parsing/merging.

    ``resolution`` is decided by the name index (design D3): file row wins
    (determinism), else a unique DB match, else a ghost target; an ambiguous
    DB match (no file row) skips the referring row with a planned-skip.
    """

    target_type: str
    name: str  # original case as written in the cell
    target_key: tuple[str, str]  # (target_type, lower(name))
    source_row_number: int  # Excel row whose cell wrote this reference
    resolution: str = LINK_TO_GHOST
    db_id: int | None = None


@dataclass
class PlannedRow:
    """A merged planned entity (one or more file rows of one type and name).

    ``fields`` carries only non-empty values of the registry column keys
    (merge: later non-empty wins — spec "Слияние одинаковых имён"; upsert:
    an empty cell therefore never reaches the update — spec "Пустая ячейка
    не затирает"). ``existing_id`` set ⇒ update plan, else create plan.
    """

    entity_type: str
    sheet: str
    key: tuple[str, str]  # (entity_type, lower(name)) — merge key
    first_row_number: int  # Excel row of the first contributing row
    fields: dict[str, Any] = field(default_factory=dict)
    links: dict[str, list[PlannedLink]] = field(default_factory=dict)
    # Date columns this row actually changed into the calendar (C3a, D7): the
    # merged field name → its visible transfer row.  A later write of the same
    # slot replaces the entry, so a superseded transfer never reaches the report.
    date_shifts: dict[str, DateShiftRow] = field(default_factory=dict)
    skipped: bool = False
    existing_id: int | None = None

    @property
    def name(self) -> str:
        return self.fields["name"]

    @property
    def is_update(self) -> bool:
        return self.existing_id is not None


@dataclass
class PlannedGhost:
    """Auto-created link target (design D5): dates are the min start /
    max end of the referring rows in the shared chronological order (design
    D2 era keys, add-era-aware-dates task 5.2 — a BC date always precedes any
    our-era one regardless of its raw ``date`` value); materialized by the
    apply pass with the matching eras."""

    entity_type: str
    key: tuple[str, str]
    name: str  # original case of the first referring cell
    min_start: GameCoord
    max_end: GameCoord
    min_start_bc: bool = False  # era of min_start (add-era-aware-dates, D7)
    max_end_bc: bool = False  # era of max_end
    referenced_by: list[tuple[str, int]] = field(default_factory=list)  # (sheet, row)


@dataclass
class ImportPlan:
    path: Path
    entities: dict[tuple[str, str], PlannedRow] = field(default_factory=dict)
    ghosts: dict[tuple[str, str], PlannedGhost] = field(default_factory=dict)
    skipped_rows: list[RowIssue] = field(default_factory=list)
    fatal_errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    # DB lower(name)s with more than one row per type (task 2.3 GROUP BY).
    ambiguous_names: dict[str, set[str]] = field(default_factory=dict)
    name_index: dict[tuple[str, str], NameIndexEntry] = field(default_factory=dict)
    # True once the DB name index was built (an empty index with built=True
    # legitimately means "none of the plan's names are in the DB").
    name_index_built: bool = False

    @property
    def has_fatal(self) -> bool:
        return bool(self.fatal_errors)

    @property
    def planned_rows(self) -> list[PlannedRow]:
        """Merged entities that will actually be created/updated."""
        return [row for row in self.entities.values() if not row.skipped]

    def iter_links(self) -> Iterator[PlannedLink]:
        for row in self.entities.values():
            if row.skipped:
                continue
            yield from (link for links in row.links.values() for link in links)

    def lookup_row(self, target_type: str, name: str) -> PlannedRow | None:
        """Planned row of ``target_type`` by (lowercased) name — link target."""
        return self.entities.get((target_type, name.strip().lower()))


def _empty_cell(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def coord_text(coord: GameCoord, is_bc: bool) -> str:
    """Coordinate display for report lines (C3a, design D5's Iso rule):
    the ISO form when the coordinate is representable as a real date, the
    codec string otherwise; BC dates carry the «… до н.э.» marker."""
    if isinstance(coord, MonthDay):
        try:
            text = date(coord.year, coord.month, coord.day).isoformat()
        except ValueError:
            text = encode_coord(coord)
    else:
        text = encode_coord(coord)
    return f"{text} до н.э." if is_bc else text


def _era_text(coord: GameCoord, is_bc: bool) -> str:
    """The report's *old* side (task C5 1.1): the coordinate exactly as it was
    parsed, rendered by the same coordinate rule as the new side — so the
    «Стандартный» preset keeps printing the plain ISO text it printed before,
    era marker («… до н.э.») included."""
    return coord_text(coord, is_bc)


def _to_coord(
    coord: GameCoord, is_bc: bool, col, sheet_title: str, row_number: int
) -> tuple[GameCoord, DateShiftRow | None]:
    """Analyze-route coordinate → the active-calendar coordinate to write.

    The parsed coordinate arrives from the parser already carried in the
    active calendar's currency (task C5 1.1) and is clamped against that
    calendar through the pure shift policy — the import never writes a
    coordinate the calendar does not contain and never shifts in silence: a
    clamp yields the visible :class:`DateShiftRow` alongside.
    """
    shifted = shift_invalid(coord, current_calendar())
    if shifted is None:
        return coord, None
    clamped, _reason = shifted
    return clamped, DateShiftRow(
        sheet=sheet_title,
        row_number=row_number,
        field=col.label,
        old=_era_text(coord, is_bc),
        new=coord_text(clamped, is_bc),
    )


def _row_to_draft(
    spec: schema.SheetSpec,
    sheet_title: str,
    row_number: int,
    resolved: schema.ResolvedHeaders,
    row: tuple,
) -> tuple[PlannedRow | None, RowIssue | None]:
    """Parse one data row into a merge candidate, or a planned-skip issue."""
    fields: dict[str, Any] = {}
    date_shifts: dict[str, DateShiftRow] = {}
    for pos in sorted(resolved.scalars):
        col = resolved.scalars[pos]
        value = row[pos] if pos < len(row) else None
        if col.key in ("start_date", "end_date"):
            parsed = schema.parse_cell_date(value)
            # A DateParse is a "coordinate XOR problem" pair (design D2): the
            # draft consumes ``coord``, while ``problem`` — raised only by the
            # custom grammar (D5) — surfaces as this row's subject caption
            # (task 2.4, texts from xlsx_report_text).
            coord = parsed.coord if parsed is not None else None
            if coord is not None:
                # C3a (D7): the parsed coordinate is clamped against the
                # active calendar BEFORE any recording route (merging, ghosts,
                # upsert) sees it — an absent day clamps with a visible row.
                coord, shift = _to_coord(coord, parsed.is_bc, col, sheet_title, row_number)
                fields[col.key] = coord
                if shift is not None:
                    date_shifts[col.key] = shift
                # The era travels next to its date under the ORM field names
                # (start_bc/end_bc — same kwargs the entity services take):
                # the merge keeps the pair atomic, later non-empty wins both.
                era_key_field = "start_bc" if col.key == "start_date" else "end_bc"
                fields[era_key_field] = parsed.is_bc
            elif col.key == "start_date":
                if parsed is not None:  # coord None + DateParse XOR ⇒ a problem
                    caption = date_problem_caption(parsed.problem.code, parsed.problem.subject)
                else:
                    caption = (
                        "пусто" if _empty_cell(value)
                        else f"значение «{value}» не является датой"
                    )
                return None, RowIssue(sheet_title, row_number, f"дата начала: {caption}")
            # An unparsable optional end date — refuse (None) or recognized
            # problem alike — is simply not applied (the spec marks only the
            # start date as a row problem).
            continue
        if col.key == "rating":
            try:
                rating = schema.parse_cell_rating(value)
            except ValueError as exc:
                return None, RowIssue(sheet_title, row_number, str(exc))
            if rating is not None:
                fields["rating"] = rating
            continue
        if _empty_cell(value):
            continue  # non-empty values only (merge/upsert overlap rule)
        fields[col.key] = str(value).strip()

    name = fields.get("name")
    if not name:
        return None, RowIssue(sheet_title, row_number, "пустое имя")

    links: dict[str, list[PlannedLink]] = {}
    for pos, link_col in resolved.links.items():
        value = row[pos] if pos < len(row) else None
        bucket = links.setdefault(link_col.target_type, [])
        seen = {link.target_key for link in bucket}
        for target_name in schema.parse_link_cell(value):
            key = (link_col.target_type, target_name.lower())
            if key in seen:
                continue  # duplicate target within one cell
            seen.add(key)
            bucket.append(
                PlannedLink(
                    target_type=link_col.target_type,
                    name=target_name,
                    target_key=key,
                    source_row_number=row_number,
                )
            )
    draft = PlannedRow(
        entity_type=spec.entity_type,
        sheet=sheet_title,
        key=(spec.entity_type, name.lower()),
        first_row_number=row_number,
        fields=fields,
        links=links,
        date_shifts=date_shifts,
    )
    return draft, None


def _parse_sheet(
    plan: ImportPlan,
    spec: schema.SheetSpec,
    sheet_title: str,
    rows: list[tuple],
) -> None:
    """Check one known sheet's header row and merge its data rows into plan."""
    if not rows or all(_empty_cell(cell) for cell in rows[0]):
        plan.fatal_errors.append(f"Лист «{sheet_title}»: нет строки заголовков.")
        return
    missing = schema.missing_required_headers(spec, rows[0])
    if missing:
        plan.fatal_errors.append(
            f"Лист «{sheet_title}»: отсутствуют обязательные колонки: {', '.join(missing)}."
        )
        return

    resolved = schema.resolve_headers(spec, rows[0])
    for row_number, raw_row in enumerate(rows[1:], start=2):
        row = tuple(None if _empty_cell(cell) else cell for cell in raw_row)
        if not any(cell is not None for cell in row):
            continue  # fully blank row (openpyxl dimension padding) — not a row
        draft, issue = _row_to_draft(spec, sheet_title, row_number, resolved, row)
        if issue is not None:
            plan.skipped_rows.append(issue)
            continue
        merged = plan.entities.get(draft.key)
        if merged is None:
            plan.entities[draft.key] = draft
            continue
        # Merge (type, lower(name)): later non-empty fields override, links
        # accumulate (order preserved, per-row dedup already applied).
        merged.fields.update(draft.fields)
        # A transferred date slot travels with its visible report row: the
        # winning contribution replaces the transfer (or clears it when the
        # new value needed no shift) — the report lists only actual transfers.
        for date_key in ("start_date", "end_date"):
            if date_key in draft.fields:
                merged.date_shifts.pop(date_key, None)
                if date_key in draft.date_shifts:
                    merged.date_shifts[date_key] = draft.date_shifts[date_key]
        for target_type, refs in draft.links.items():
            existing = merged.links.setdefault(target_type, [])
            seen = {link.target_key for link in existing}
            existing.extend(link for link in refs if link.target_key not in seen)


def analyze_workbook(path: Path) -> ImportPlan:
    """One-shot read of the workbook (task 2.1): known sheets are parsed
    and merged, unknown ones only warned about; every file-level defect
    lands in ``plan.fatal_errors`` and aborts the rest of the analysis."""
    plan = ImportPlan(path=path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of format errors
        plan.fatal_errors.append(f"Файл повреждён или не является .xlsx: {exc}")
        return plan

    try:
        known: list[tuple[schema.SheetSpec, str]] = []
        for title in wb.sheetnames:
            spec = schema.find_sheet_by_name(title)
            if spec is None:
                plan.warnings.append(
                    f"Неизвестный лист «{title}» не входит в набор листов "
                    "импорта и импортирован не будет."
                )
            else:
                known.append((spec, title))

        by_registry_name: dict[str, list[str]] = {}
        for spec, title in known:
            by_registry_name.setdefault(
                schema.normalize_header(spec.sheet_name), []
            ).append(title)
        duplicates = [titles for titles in by_registry_name.values() if len(titles) > 1]
        if duplicates:
            for titles in duplicates:
                plan.fatal_errors.append(
                    "Дубликаты листа в нескольких написаниях: "
                    + ", ".join(f"«{t}»" for t in titles)
                )
            return plan

        if not known:
            expected = ", ".join(s.sheet_name for s in schema.all_sheets())
            plan.fatal_errors.append(
                f"В файле нет ни одного знакомого листа. Ожидаемые листы: {expected}."
            )
            return plan

        for spec, title in known:
            _parse_sheet(plan, spec, title, list(wb[title].iter_rows(values_only=True)))

        if not plan.fatal_errors and not plan.entities and not plan.skipped_rows:
            plan.fatal_errors.append(
                "Файл пустой: знакомые листы не содержат ни одной строки данных."
            )
    finally:
        wb.close()
    return plan


def _collect_wanted_names(plan: ImportPlan) -> dict[str, set[str]]:
    """lower(name)s per type that the plan needs resolved against the DB."""
    wanted: dict[str, set[str]] = {}
    for row in plan.entities.values():
        wanted.setdefault(row.entity_type, set()).add(row.key[1])
    for row in plan.entities.values():
        for refs in row.links.values():
            for ref in refs:
                wanted.setdefault(ref.target_type, set()).add(ref.target_key[1])
    return wanted


async def build_name_index(plan: ImportPlan, session: AsyncSession) -> None:
    """Name index of design D3 — DB queries live in the repository layer
    (``query_lower_name_index``, audit B6); here the raw occurrences become
    plan structures. An update target is planned only for a unique
    match; ambiguous row names stay create-plans (existing_id is None)."""
    wanted = _collect_wanted_names(plan)
    for entity_type, names in wanted.items():
        ambiguous, by_lower = await query_lower_name_index(session, entity_type, names)
        if ambiguous:
            plan.ambiguous_names.setdefault(entity_type, set()).update(ambiguous)
        for lower_key in set(by_lower) | ambiguous:
            plan.name_index[(entity_type, lower_key)] = NameIndexEntry(
                models=list(by_lower.get(lower_key, []))
            )
    for row in plan.entities.values():
        entry = plan.name_index.get(row.key)
        if entry is not None and entry.is_unique:
            row.existing_id = entry.models[0].id
    plan.name_index_built = True


def _grow_ghost(plan: ImportPlan, ref: PlannedLink, row: PlannedRow) -> None:
    """Expand the ghost buffer with one reference (design D5: min start /
    max end of the referring rows; a missing end date counts as its start).
    Bounds compare through the shared era key (design D2/D7): raw ``date``
    objects would order BC years backwards and mix the eras wrongly."""
    start = row.fields["start_date"]
    start_bc = bool(row.fields.get("start_bc"))
    end = row.fields.get("end_date") or start
    end_bc = bool(row.fields.get("end_bc")) if row.fields.get("end_date") else start_bc
    start_key = era_key(start, start_bc)
    end_key = era_key(end, end_bc)
    ghost = plan.ghosts.get(ref.target_key)
    if ghost is None:
        plan.ghosts[ref.target_key] = PlannedGhost(
            entity_type=ref.target_type,
            key=ref.target_key,
            name=ref.name,  # first reference keeps its case
            min_start=start,
            max_end=end,
            min_start_bc=start_bc,
            max_end_bc=end_bc,
            referenced_by=[(row.sheet, ref.source_row_number)],
        )
        return
    if start_key < era_key(ghost.min_start, ghost.min_start_bc):
        ghost.min_start, ghost.min_start_bc = start, start_bc
    if end_key > era_key(ghost.max_end, ghost.max_end_bc):
        ghost.max_end, ghost.max_end_bc = end, end_bc
    source = (row.sheet, ref.source_row_number)
    if source not in ghost.referenced_by:
        ghost.referenced_by.append(source)


def _resolve_row_links(plan: ImportPlan, row: PlannedRow) -> RowIssue | None:
    """Classify one row's link references; None if all resolve fine.

    A reference to an ambiguous DB name that no file row defines makes the
    whole referring row a planned skip (task 2.3); a name absent from both
    file and DB becomes a ghost target.
    """
    for refs in row.links.values():
        for ref in refs:
            target = plan.entities.get(ref.target_key)
            if target is not None and not target.skipped:
                ref.resolution, ref.db_id = LINK_TO_FILE, None
                continue
            entry = plan.name_index.get(ref.target_key)
            if entry is not None and entry.is_unique:
                ref.resolution, ref.db_id = LINK_TO_DB, entry.models[0].id
                continue
            if entry is not None and entry.is_ambiguous:
                label = schema.LINK_COLUMNS_BY_TARGET[ref.target_type].label
                sheet_name = schema.sheet_for(ref.target_type).sheet_name
                return RowIssue(
                    row.sheet,
                    ref.source_row_number,
                    ambiguous_link_reference(label, ref.name, sheet_name),
                )
            # Unknown target → ghost, unless the name index was never built
            # (no session *and* the file defines no rows of the target type):
            # then membership in the DB cannot be told apart from absence.
            has_file_rows_of_type = any(
                other.entity_type == ref.target_type
                and not other.skipped
                for other in plan.entities.values()
            )
            if not plan.name_index_built and not has_file_rows_of_type:
                plan.fatal_errors.append(
                    f"Лист «{row.sheet}», строка {ref.source_row_number}: ссылка "
                    f"«{ref.name}» не может быть проверена — индекс имён не построен"
                )
            ref.resolution, ref.db_id = LINK_TO_GHOST, None
    return None


def resolve_links(plan: ImportPlan) -> None:
    """Resolve every link reference, iterating to a fixpoint: skipping a row
    can remove it as a file-resolvable target and cascade its referrers to
    the DB/ghost/ambiguous branches (the skip set only grows, so this ends).
    Ghosts are finally grown from surviving rows only."""
    while True:
        new_skips: list[RowIssue] = []
        for row in list(plan.entities.values()):
            if row.skipped:
                continue
            issue = _resolve_row_links(plan, row)
            if issue is not None:
                row.skipped = True
                new_skips.append(issue)
        plan.skipped_rows.extend(new_skips)
        if not new_skips:
            break
    plan.ghosts = {}
    for row in plan.entities.values():
        if row.skipped:
            continue
        for refs in row.links.values():
            for ref in refs:
                if ref.resolution == LINK_TO_GHOST:
                    _grow_ghost(plan, ref, row)
